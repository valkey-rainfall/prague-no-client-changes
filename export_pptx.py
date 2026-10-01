#!/usr/bin/env python3
"""PowerPoint backup of the deck: prague-deck.pptx, 16:9, one full-bleed PNG per body slide, with a still of the
lit sign-on (wordmark, title, name) as the first slide in place of the live intro.

The web deck (prague-deck/index.html) is the talk. This file exists for the cases a browser cannot cover: a
venue that insists on a .pptx, a laptop without Chrome, or someone who wants to read the slides in the tool
they know. It carries no sound and no animation; the intro's four loader frames and the CRT-off are dropped,
since in PowerPoint they would be four clicks of still pictures.

usage: export_pptx.py [out.pptx] [--signon PNG | --no-signon]
  The sign-on still is assets/signon-lit.png, the 07-lit.png frame test_deck.mjs captures (copy a fresh one
  there after changing the intro). Needs python-pptx (pip install python-pptx) and prague-deck/slides/.
"""
import json
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
DECK = HERE / 'prague-deck'
W, H = 12192000, 6858000                      # 13.333 x 7.5 in, PowerPoint's own 16:9


def main(argv):
    out = next((Path(a) for a in argv if a.endswith('.pptx')), HERE / 'prague-deck.pptx')
    signon = None if '--no-signon' in argv else (Path(argv[argv.index('--signon') + 1]) if '--signon' in argv else HERE / 'assets' / 'signon-lit.png')
    slides = json.loads((DECK / 'slides' / 'slides.json').read_text())

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    blank = prs.slide_layouts[6]

    def add(png, notes=None):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(png), 0, 0, width=Emu(W), height=Emu(H))   # every PNG is 16:9, so no letterboxing
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    if signon:
        add(signon, 'Still of the live sign-on. The web deck (prague-deck/index.html) plays the real intro, with sound.')
    for rel in slides:
        add(DECK / rel)
    prs.save(out)
    n = len(slides) + bool(signon)
    print(f'{out} : {n} slides ({"sign-on still + " if signon else ""}{len(slides)} body), {out.stat().st_size / 1e6:.1f} MB'
          )


if __name__ == '__main__':
    main(sys.argv[1:])
